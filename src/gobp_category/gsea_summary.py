#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GSEA summary table generator for the GO BP pipeline
======================================================

This script reads a GSEA result CSV, classifies it through
gsea_classification.py, and writes a formatted Excel summary table.

By default, summaries are grouped by Category_Merged. Use
--category-col Category to summarize by the original fine-grained English
Category instead.

Usage
-----
    python -m gobp_pipeline.gsea_summary input.csv output.xlsx
    python -m gobp_pipeline.gsea_summary input.csv output.xlsx --category-col Category
"""

import argparse
import re
from collections import Counter

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

if __package__:
    from . import gsea_classification as gsea
    from .workbook_names import unique_sheet_title
else:
    import sys
    from pathlib import Path

    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from gobp_pipeline import gsea_classification as gsea
    from gobp_pipeline.workbook_names import unique_sheet_title


# Display order for original fine-grained categories.
ORIGINAL_CATEGORY_ORDER = []
for _cat_cn, _ in gsea.RULES:
    _cat_en = gsea._CATEGORY_EN.get(_cat_cn)
    if _cat_en and _cat_en not in ORIGINAL_CATEGORY_ORDER:
        ORIGINAL_CATEGORY_ORDER.append(_cat_en)


# Display order for merged categories. Categories not listed here but present in
# data are appended after these, alphabetically.
MERGED_CATEGORY_ORDER = [
    "Mitochondrial Function & OXPHOS",
    "DNA Replication, Repair & Genome Stability",
    "Translation & Ribosome Biogenesis",
    "RNA Processing & Splicing",
    "Transcription & Chromatin",
    "General Metabolism",
    "Ion Transport & Homeostasis",
    "Protein Folding & Degradation",
    "Extracellular Matrix Organization",
    "Chemotaxis, Migration, Adhesion & Cytoskeleton",
    "T Cell Activation & Signaling",
    "B Cell & Antibody",
    "Leukocyte Activation & Migration",
    "Cytokine & Interleukin Response",
    "Antigen Processing & Presentation",
    "Cytotoxicity & Cell Killing",
    "Immune Effector & Tolerance",
    "Apoptosis, Programmed Cell Death & Autophagy",
    "Interferon & Antiviral Response",
    "Stress Response",
    "Wound Healing & Regeneration",
    "Cell Cycle Regulation & Proliferation",
    "Growth Factor/RTK Signaling",
    "G Protein & MAPK Cascade Signaling",
    "Morphogen Signaling",
    "Intracellular Signal Transduction (General)",
    "Vesicle Transport & Secretion",
    "Development, neural and organogenesis",
    "Other/Uncategorized",
]


CATEGORY_CN_BY_EN = {en: cn for cn, en in gsea._CATEGORY_EN.items()}

MERGED_CATEGORY_CN = {
    "DNA Replication, Repair & Genome Stability": "DNA复制、修复与基因组稳定",
    "Chemotaxis, Migration, Adhesion & Cytoskeleton": "趋化、迁移、粘附与细胞骨架",
    "General Metabolism": "一般代谢",
    "Interferon & Antiviral Response": "干扰素与抗病毒反应",
    "Development, neural and organogenesis": "发育、神经与器官发生",
}


SUBDESC_BY_CN = {
    "线粒体功能与OXPHOS": "氧化磷酸化、ATP合成、电子传递链、线粒体翻译、铁硫簇",
    "DNA复制": "DNA复制起始、复制叉、解旋酶、DNA合成与拓扑变化",
    "DNA损伤应答与修复": "DNA损伤应答、双链断裂修复、同源重组、端粒维持、交叉链接修复",
    "有丝分裂与染色体分离": "有丝分裂、纺锤体组装、着丝粒/动粒、染色体分离、胞质分裂",
    "翻译与核糖体生物发生": "翻译、核糖体生物发生、肽链合成、核糖核蛋白复合体组装",
    "RNA加工与剪接": "mRNA加工、剪接体、RNA修饰/编辑、核输出、polyadenylation",
    "转录与染色质": "转录调控、染色质重塑、组蛋白修饰、DNA甲基化、基因表达调控",
    "脂质与固醇代谢": "脂质代谢、脂肪酸氧化、胆固醇/固醇合成、磷脂、鞘脂、脂蛋白",
    "糖代谢与糖酵解": "糖酵解、糖异生、葡萄糖代谢、磷酸戊糖途径、糖原代谢",
    "氨基酸代谢": "氨基酸代谢、转氨基、尿素循环、谷氨酰胺/谷胱甘肽代谢",
    "核苷酸与嘌呤代谢": "核苷酸合成、嘌呤/嘧啶代谢、核苷碱基、GMP/IMP/AMP生物合成",
    "代谢（通用/其他）": "小分子、有机酸、磷酸盐、辅酶及其他代谢过程",
    "离子转运与稳态": "离子跨膜转运、金属离子稳态、pH稳态、渗透压调节、维生素/辅因子",
    "蛋白质折叠与降解": "蛋白折叠、伴侣蛋白、泛素-蛋白酶体、蛋白降解、ER加工与糖基化",
    "蛋白质定位与靶向": "蛋白质定向、细胞内定位、膜和细胞器导入导出",
    "蛋白质修饰与加工": "翻译后修饰、磷酸化、乙酰化、糖基化与蛋白成熟",
    "蛋白质复合体组装与组织": "蛋白质复合体的组装、拆解与重塑",
    "小分子与跨膜转运": "小分子、代谢物和药物的跨膜转运与摄取",
    "细胞器与膜组织": "细胞器、核膜及其他膜结构的组装与维护",
    "细胞壁组织与生物发生": "植物和微生物细胞壁的组装、修饰与生物发生",
    "光合作用与叶绿体功能": "光合作用、光系统、叶绿体和类囊体功能",
    "植物发育与生理": "花、种子、分生组织、维管组织和植物发育过程",
    "微生物生长与生物膜": "生物膜、孢子形成、菌丝生长与菌毛装配",
    "宿主-微生物相互作用": "微生物对宿主过程的调节及共生相互作用",
    "解毒与异生物清除": "化学物、异生物和有毒代谢物的解毒与清除",
    "病毒生命周期": "病毒进入、复制、组装、出芽与释放",
    "刺激与激素响应（其他）": "化学、物理及激素刺激的细胞响应",
    "细胞外基质组织": "细胞外基质、胶原、基底膜、蛋白聚糖与外部包裹结构",
    "细胞黏附": "细胞黏附、钙黏蛋白、整合素、细胞-细胞连接、灶性黏附",
    "细胞迁移与趋化": "细胞迁移、趋化、运动性、归巢、趋化因子响应",
    "细胞骨架组织、连接与极性": "细胞骨架组织、肌动蛋白/微管、细胞连接、极性、纤毛/细胞突起",
    "T细胞激活与信号": "T细胞激活、TCR信号、CD4/CD8 T细胞、T细胞增殖与分化",
    "B细胞与抗体": "B细胞激活/增殖、免疫球蛋白、类别转换、BCR信号、抗体应答",
    "白细胞激活与迁移": "白细胞激活、髓系细胞、中性粒细胞、巨噬细胞、吞噬、树突状细胞",
    "细胞因子与白细胞介素响应": "细胞因子产生/响应、白细胞介素、JAK-STAT信号、淋巴因子",
    "抗原加工与递呈": "抗原加工、抗原递呈、MHC I/II类分子、肽抗原结合",
    "杀伤功能": "细胞杀伤、白细胞介导免疫、细胞毒性、靶细胞杀伤调控",
    "免疫效应与耐受": "免疫应答、NK细胞、体液免疫、补体、炎症、抗菌防御、免疫耐受",
    "凋亡、程序性死亡与自噬": "凋亡、程序性细胞死亡、坏死性凋亡、焦亡、自噬/巨自噬",
    "干扰素与抗病毒反应": "干扰素产生与响应、RIG-I/MDA-5 信号和抗病毒防御",
    "应激反应": "氧化应激、化学/缺氧/损伤及营养压力响应",
    "细胞周期调控与增殖": "细胞周期、G1/S和G2/M转换、检查点、增殖、细胞生长",
    "生长因子/RTK信号": "生长因子信号、ErbB/VEGF/FGF/EGF/PDGF、受体酪氨酸激酶",
    "G蛋白与MAPK级联信号": "GPCR、GTPase、MAPK级联(ERK/JNK/p38)、Ras/Rho/Rac/Cdc42",
    "形态发生素信号": "Notch、Wnt/beta-catenin、BMP、Hedgehog、TGF-beta、Eph/ephrin",
    "细胞内信号转导(通用)": "信号转导、激酶/磷酸化、级联放大、受体、激素、细胞间信号",
    "囊泡运输与分泌": "囊泡运输、胞吐/分泌、内体/溶酶体、高尔基体、内吞、蛋白定位",
    "神经系统发育": "神经发生、轴突/树突、胶质细胞、髓鞘、突触组织与神经传递",
    "心脏发育": "心脏、心肌、心室/心房、心内膜和瓣膜相关发育",
    "血管与内皮发育": "血管生成、内皮发育、血管形态发生、凝血和循环调节",
    "骨骼与软骨发育": "骨、软骨、骨化、成骨细胞和软骨细胞分化",
    "肌肉发育": "肌肉器官、肌细胞、肌原纤维、肌节和肌肉收缩",
    "上皮与皮肤发育": "上皮、表皮、皮肤屏障、毛囊和角化过程",
    "器官发育": "肾、眼、耳、肺、肝、胰腺、消化/泌尿生殖等器官发育",
    "胚胎与体轴发育": "胚胎、原肠胚、胚层、体轴、模式形成和分节",
    "形态发生": "解剖结构形态发生、分支、内陷、管腔发育和组织形态变化",
    "细胞分化与命运决定": "细胞分化、细胞命运、谱系承诺、干细胞和成熟过程",
    "伤口愈合与再生": "伤口愈合、组织重塑和再生",
    "系统生理与行为": "行为、体液、血压、肾小球滤过、生殖、节律和系统过程",
    "发育·其他（去分化背景项）": "未归入具体背景子类的发育、形成和区域化过程",
    "其他/未归类": "未归入上述功能分类的其他生物学过程",
}

SUBDESC_BY_EN = {
    gsea._CATEGORY_EN[cn]: desc
    for cn, desc in SUBDESC_BY_CN.items()
    if cn in gsea._CATEGORY_EN
}

MERGED_SUBDESC = {
    "DNA Replication, Repair & Genome Stability": "DNA复制、损伤修复、有丝分裂、染色体分离与基因组稳定性",
    "Chemotaxis, Migration, Adhesion & Cytoskeleton": "细胞趋化、迁移、黏附、连接、极性和细胞骨架重塑",
    "General Metabolism": "脂质、糖、氨基酸、核苷酸和其他一般代谢过程",
    "Interferon & Antiviral Response": "干扰素产生与响应、RIG-I/MDA-5 信号和抗病毒防御",
    "Development, neural and organogenesis": "发育、神经系统、器官发生、形态发生、细胞命运和系统生理背景项",
}


def validate_required_columns(df: pd.DataFrame):
    gsea.validate_gsea_result_columns(df)


def prepare_classified_gsea_results(df: pd.DataFrame) -> pd.DataFrame:
    """Apply the shared GO BP classifier before calculating Summary statistics."""
    return gsea.classify_gsea_results(df)


def detect_adjacent_comparisons(comparisons: list) -> tuple:
    """Detect adjacent and fallback comparison groups from labels like R3_vs_R1."""
    adjacent = []
    fallback = []
    label = {}

    round_nums = set()
    comp_pairs = {}
    for comp in sorted(comparisons):
        m = re.match(r"R(\d+)_vs_R(\d+)", str(comp))
        if m:
            x, y = int(m.group(1)), int(m.group(2))
            comp_pairs[comp] = (x, y)
            round_nums.add(x)
            round_nums.add(y)
            label[comp] = f"R{x} vs R{y}"
        else:
            label[comp] = str(comp).replace("_vs_", " vs ")
            fallback.append(comp)

    sorted_rounds = sorted(round_nums)
    for comp in sorted(comp_pairs.keys()):
        x, y = comp_pairs[comp]
        idx_x = sorted_rounds.index(x)
        idx_y = sorted_rounds.index(y)
        if abs(idx_x - idx_y) == 1:
            adjacent.append(comp)
        else:
            fallback.append(comp)

    return adjacent, fallback, label


def select_representative_comparison(cat_df: pd.DataFrame, adjacent: list, fallback: list, label: dict):
    """Choose the strongest comparison for one category."""
    best_info = None
    best_score = -1

    for comp_set in (adjacent, fallback):
        if best_info is not None:
            break
        for comp in comp_set:
            sub = cat_df[cat_df["Comparison"] == comp]
            if len(sub) == 0:
                continue
            med_abs = sub["absNES"].median()
            score = len(sub) * med_abs
            if score > best_score:
                best_score = score
                dirs = gsea.normalize_gsea_directions(sub).value_counts()
                dom_dir = "上调" if dirs.get("上调", 0) >= dirs.get("下调", 0) else "下调"
                dir_en = "Up" if dom_dir == "上调" else "Down"
                best_info = {
                    "comparison": comp,
                    "label": f"{label[comp]}({dir_en})",
                    "direction": dom_dir,
                    "direction_en": dir_en,
                    "n_pathways": len(sub),
                    "median_absNES": round(med_abs, 2),
                    "max_absNES": round(sub["absNES"].max(), 2),
                }

    return best_info


def extract_core_genes(top_df: pd.DataFrame, max_genes: int = 20) -> tuple:
    """Extract recurrent core genes from the representative pathways."""
    if "core_enrichment" not in top_df.columns:
        return "", 0

    gene_counter = Counter()
    gene_max_absnes = {}

    for _, row in top_df.iterrows():
        if pd.isna(row["core_enrichment"]):
            continue
        genes = str(row["core_enrichment"]).split("/")
        for gene in genes:
            gene = gene.strip()
            if not gene:
                continue
            gene_counter[gene] += 1
            if gene not in gene_max_absnes or row["absNES"] > gene_max_absnes[gene]:
                gene_max_absnes[gene] = row["absNES"]

    multi = [(g, c, gene_max_absnes[g]) for g, c in gene_counter.items() if c >= 2]
    multi.sort(key=lambda x: (-x[1], -x[2], x[0]))

    single = [(g, c, gene_max_absnes[g]) for g, c in gene_counter.items() if c == 1]
    single.sort(key=lambda x: (-x[2], x[0]))

    selected = (multi + single)[:max_genes]
    gene_list = [g for g, _, _ in selected]
    return ", ".join(f'"{g}"' for g in gene_list), len(gene_list)


def ordered_categories(df: pd.DataFrame, category_col: str) -> list:
    present = set(df[category_col].dropna().unique())
    base_order = MERGED_CATEGORY_ORDER if category_col == "Category_Merged" else ORIGINAL_CATEGORY_ORDER
    ordered = [cat for cat in base_order if cat in present]
    ordered.extend(sorted(present - set(ordered)))
    return ordered


def describe_summary_category(cat: str, cat_df: pd.DataFrame, category_col: str) -> tuple:
    """Return (Chinese label, English label, Chinese subdescription)."""
    if category_col == "Category_Merged":
        cat_cn = MERGED_CATEGORY_CN.get(cat)
        if cat_cn is None and "Category_CN" in cat_df.columns and cat_df["Category_CN"].nunique() == 1:
            cat_cn = cat_df["Category_CN"].iloc[0]
        cat_cn = cat_cn or CATEGORY_CN_BY_EN.get(cat, cat)
        subdesc = MERGED_SUBDESC.get(cat) or SUBDESC_BY_EN.get(cat, "")
        return cat_cn, cat, subdesc

    cat_cn = CATEGORY_CN_BY_EN.get(cat)
    if cat_cn is None and "Category_CN" in cat_df.columns and len(cat_df) > 0:
        mode = cat_df["Category_CN"].mode(dropna=True)
        cat_cn = mode.iloc[0] if len(mode) else cat
    cat_cn = cat_cn or cat
    return cat_cn, cat, SUBDESC_BY_EN.get(cat, "")


def _write_gsea_summary_sheet(ws, results: list):
    """Populate one Group sheet using the existing Summary layout."""
    header_fill = PatternFill(start_color="FF0E2841", end_color="FF0E2841", fill_type="solid")
    num_fill = PatternFill(start_color="FF0E2841", end_color="FF0E2841", fill_type="solid")
    white = "FFFFFFFF"
    dark_blue = "FF173F5F"
    light_blue = "FF00B0F0"

    font_header = Font(name="微软雅黑", size=10, bold=True, color=white)
    font_num = Font(name="Arial", size=9, bold=True, color=white)
    font_catname = Font(name="微软雅黑", size=9, bold=True, color=dark_blue)
    font_subdesc = Font(name="微软雅黑", size=8, bold=False, color=dark_blue)
    font_en = Font(name="Arial", size=8, bold=False, color=dark_blue)
    font_pathway = Font(name="Arial", size=7, bold=False, color=dark_blue)
    font_gene = Font(name="Arial", size=7, bold=False, color=dark_blue)
    font_round = Font(name="Arial", size=9, bold=True, color=light_blue)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    medium = Side(style="medium", color="FF000000")
    border_all = Border(top=medium, bottom=medium, left=medium, right=medium)

    headers = ["序号", "生物学功能类型", "English Category", "代表通路", "核心差异基因", "基因差异轮次"]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_all
    ws.row_dimensions[1].height = 28

    for i, row in enumerate(results):
        r1 = 2 + i * 2
        r2 = r1 + 1

        ws.cell(row=r1, column=1, value=i + 1)
        ws.merge_cells(start_row=r1, start_column=1, end_row=r2, end_column=1)

        ws.cell(row=r1, column=2, value=row["category_cn"])
        ws.cell(row=r2, column=2, value=row["subdesc"])

        ws.cell(row=r1, column=3, value=row["category_en"])
        ws.merge_cells(start_row=r1, start_column=3, end_row=r2, end_column=3)

        ws.cell(row=r1, column=4, value=row["pathways"])
        ws.merge_cells(start_row=r1, start_column=4, end_row=r2, end_column=4)

        ws.cell(row=r1, column=5, value=row["genes"])
        ws.merge_cells(start_row=r1, start_column=5, end_row=r2, end_column=5)

        ws.cell(row=r1, column=6, value=row["round_label"])
        ws.merge_cells(start_row=r1, start_column=6, end_row=r2, end_column=6)

        for r in [r1, r2]:
            cell = ws.cell(row=r, column=1)
            cell.font = font_num
            cell.fill = num_fill
            cell.alignment = align_center
            cell.border = border_all

            cell = ws.cell(row=r, column=3)
            cell.font = font_en
            cell.alignment = align_center
            cell.border = border_all

            cell = ws.cell(row=r, column=4)
            cell.font = font_pathway
            cell.alignment = align_left
            cell.border = border_all

            cell = ws.cell(row=r, column=5)
            cell.font = font_gene
            cell.alignment = align_left
            cell.border = border_all

            cell = ws.cell(row=r, column=6)
            cell.font = font_round
            cell.alignment = align_center
            cell.border = border_all

        cell_top = ws.cell(row=r1, column=2)
        cell_top.font = font_catname
        cell_top.alignment = align_center
        cell_top.border = Border(top=medium, bottom=None, left=medium, right=medium)

        cell_bottom = ws.cell(row=r2, column=2)
        cell_bottom.font = font_subdesc
        cell_bottom.alignment = align_center
        cell_bottom.border = Border(top=None, bottom=medium, left=medium, right=medium)

        ws.row_dimensions[r1].height = 30
        ws.row_dimensions[r2].height = 30

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 52
    ws.column_dimensions["E"].width = 55
    ws.column_dimensions["F"].width = 16

    ws.freeze_panes = "A2"


def write_gsea_summary_workbook(results_by_group, out_path: str):
    """Write one Summary sheet per Group; a list retains the legacy one-sheet API."""
    if isinstance(results_by_group, list):
        results_by_group = {"GSEA_category_summary": results_by_group}
    if not results_by_group:
        raise ValueError("No Group data available for Summary output.")

    wb = openpyxl.Workbook()
    used_titles = set()
    for index, (group, results) in enumerate(results_by_group.items()):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = unique_sheet_title(group, used_titles)
        _write_gsea_summary_sheet(ws, results)
    wb.save(out_path)


def _summarize_group(df: pd.DataFrame, category_col: str, top_pathways: int, top_genes: int) -> list:
    """Select representative comparisons independently within one Group."""
    comparisons = sorted(df["Comparison"].dropna().unique())
    adjacent, fallback, label = detect_adjacent_comparisons(comparisons)
    results = []
    for cat in ordered_categories(df, category_col):
        cat_df = df[df[category_col] == cat]
        if len(cat_df) == 0:
            continue

        info = select_representative_comparison(cat_df, adjacent, fallback, label)
        if info is None:
            print(f"  跳过: {cat}（无有效轮次）")
            continue

        selected = cat_df[cat_df["Comparison"] == info["comparison"]].sort_values("absNES", ascending=False)
        top = selected.head(top_pathways)
        pathways = ", ".join(top["Description"].tolist())
        genes_str, n_genes = extract_core_genes(top, max_genes=top_genes)
        category_cn, category_en, subdesc = describe_summary_category(cat, cat_df, category_col)

        results.append({
            "category_cn": category_cn,
            "category_en": category_en,
            "subdesc": subdesc,
            "pathways": pathways,
            "genes": genes_str,
            "n_genes": n_genes,
            "round_label": info["label"],
            "comparison": info["comparison"],
            "direction": info["direction"],
            "n_pathways_total": info["n_pathways"],
            "n_pathways_top": len(top),
            "median_absNES": info["median_absNES"],
            "max_absNES": info["max_absNES"],
        })

    return results


def generate_gsea_summary(
    input_path: str,
    output_path: str,
    top_pathways: int = 5,
    top_genes: int = 20,
    category_col: str = "Category_Merged",
):
    """Classify GSEA terms and write one Summary sheet for each input Group."""
    df = prepare_classified_gsea_results(pd.read_csv(input_path, encoding="utf-8-sig"))
    if df.empty:
        raise ValueError("GSEA input contains no rows.")
    if category_col not in df.columns:
        raise ValueError(f"Unknown category column: {category_col}")

    has_group = "Group" in df.columns
    if has_group:
        names = df["Group"].astype("string").str.strip()
        if names.isna().any() or names.eq("").any():
            raise ValueError("Group contains missing or blank values.")
        df = df.assign(Group=names)
        groups = df.groupby("Group", sort=False)
    else:
        groups = [("GSEA_category_summary", df)]

    print(f"Hierarchy available: {gsea._HIERARCHY_AVAILABLE}")
    print(f"已加载 {len(df)} 行，按 {category_col} 分组汇总")
    results_by_group = {}
    for group_name, group_df in groups:
        results = _summarize_group(group_df, category_col, top_pathways, top_genes)
        results_by_group[str(group_name)] = results
        print(f"  {group_name}: {len(group_df)} 条通路，{len(results)} 个分类")

    write_gsea_summary_workbook(results_by_group, output_path)
    print(f"\n已保存: {output_path}")

    wb = openpyxl.load_workbook(output_path, read_only=True)
    print(f"验证: {len(wb.sheetnames)} 个 Group 工作表: {', '.join(wb.sheetnames)}")
    wb.close()

    return results_by_group if has_group else results_by_group["GSEA_category_summary"]


def main(argv=None):
    """Command-line entry point for the GSEA Summary workbook."""
    parser = argparse.ArgumentParser(
        description="GSEA summary table generator using the shared GO BP classifier",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python -m gobp_pipeline.gsea_summary GSEA_HS_fdr0.05.csv summary.xlsx
  python -m gobp_pipeline.gsea_summary GSEA_HS_fdr0.05.csv summary_fine.xlsx --category-col Category
        """,
    )
    parser.add_argument("input", help="GSEA result CSV path")
    parser.add_argument("output", help="Output Excel path")
    parser.add_argument("--top-pathways", type=int, default=5, help="Representative pathways per category")
    parser.add_argument("--top-genes", type=int, default=20, help="Core genes per category")
    parser.add_argument(
        "--category-col",
        choices=["Category", "Category_Merged"],
        default="Category_Merged",
        help="Summary category column. Default: Category_Merged",
    )

    args = parser.parse_args(argv)
    generate_gsea_summary(
        args.input,
        args.output,
        top_pathways=args.top_pathways,
        top_genes=args.top_genes,
        category_col=args.category_col,
    )


# Historical entry points remain available for existing notebooks.
ensure_classified = prepare_classified_gsea_results
best_round = select_representative_comparison
category_display = describe_summary_category
build_xlsx = write_gsea_summary_workbook
generate_summary_table = generate_gsea_summary


if __name__ == "__main__":
    main()
