"""Export per-Group GSEA category statistics by comparison and direction."""

import argparse
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

if __package__:
    from .gsea_classification import normalize_gsea_directions, summarize_gsea_groups
    from .workbook_names import unique_group_filename, unique_sheet_title
else:
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from gobp_pipeline.gsea_classification import normalize_gsea_directions, summarize_gsea_groups
    from gobp_pipeline.workbook_names import unique_group_filename, unique_sheet_title


STAT_COLUMNS = (
    "Category", "n", "median_absNES", "max_absNES", "min_padj",
    "top_term", "top_NES", "prominence",
)


def _write_comparison_sheet(sheet, statistics: pd.DataFrame, category_col: str) -> None:
    """Write the eight-column layout of the existing comparison workbook."""
    sheet.append([category_col, *STAT_COLUMNS[1:]])
    for values in statistics.itertuples(index=False, name=None):
        sheet.append([None if pd.isna(value) else value for value in values])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="173B54")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
    for column, width in {
        "A": 48, "B": 10, "C": 17, "D": 16,
        "E": 16, "F": 65, "G": 14, "H": 16,
    }.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=8):
        for cell in row:
            if cell.column == 5:
                cell.number_format = "0.00E+00"
            elif cell.column != 6:
                cell.number_format = "0.00"


def generate_group_comparison_workbooks(
    input_path: str, output_dir: str, category_col: str = "Category"
) -> dict[str, Path]:
    """Create one XLSX per Group, with a sheet for each Comparison and direction."""
    if category_col not in {"Category", "Category_Merged"}:
        raise ValueError("category_col must be Category or Category_Merged.")
    df = pd.read_csv(input_path, encoding="utf-8-sig")
    required = {"Group", "Comparison", "Description", "NES", "p.adjust", category_col}
    missing = sorted(required - set(df.columns))
    if missing:
        raise ValueError(f"Classified GSEA CSV is missing: {', '.join(missing)}")
    if df.empty:
        raise ValueError("Classified GSEA CSV contains no rows.")
    normalize_gsea_directions(df)
    grouped_statistics = summarize_gsea_groups(df, category_col=category_col)

    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    used_filenames = set()
    written = {}
    for group, aggregates in grouped_statistics.items():
        workbook = openpyxl.Workbook()
        used_titles = set()
        for index, ((comparison, direction), statistics) in enumerate(aggregates.items()):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            suffix = "up" if direction == "上调" else "down"
            sheet.title = unique_sheet_title(f"{comparison}_{suffix}", used_titles)
            _write_comparison_sheet(sheet, statistics, category_col)
        filename = unique_group_filename(group, used_filenames)
        target = output_path / f"GSEA_category_summary_{filename}.xlsx"
        workbook.save(target)
        written[str(group)] = target
    return written


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description="Export one comparison-category workbook per GSEA Group")
    parser.add_argument("input", help="Classified GSEA CSV with a Group column")
    parser.add_argument("output_dir", help="Directory for one XLSX per Group")
    parser.add_argument(
        "--category-col", choices=["Category", "Category_Merged"], default="Category",
        help="Category level for statistics (default: Category)",
    )
    args = parser.parse_args(argv)
    for group, path in generate_group_comparison_workbooks(
        args.input, args.output_dir, category_col=args.category_col
    ).items():
        print(f"{group}: {path}")


if __name__ == "__main__":
    main()
